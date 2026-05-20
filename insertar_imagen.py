"""
Inserta una imagen en el placeholder de una figura específica del Word.
Uso:
    python insertar_imagen.py <num_figura> <ruta_imagen>
"""
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

RUTA_DOC = Path("/Users/dev/Downloads/BASESDEDATOS2/laboratorio18abrilsegundo/LAB_18_ABRIL_Joan_Mateo.docx")


def insertar_en_figura(num_figura: int, ruta_imagen: Path) -> bool:
    doc = Document(str(RUTA_DOC))
    body = doc.element.body
    target_text = f"Figura {num_figura}."

    children = list(body.iterchildren())
    for idx, el in enumerate(children):
        # Buscamos el párrafo que contiene "Figura N."
        if el.tag.endswith('}p') and target_text in (el.text or ''):
            # Retrocedemos para encontrar la tabla previa (la del placeholder)
            for back in range(idx - 1, -1, -1):
                prev = children[back]
                if prev.tag.endswith('}tbl'):
                    # Obtenemos la tabla python-docx correspondiente
                    for tbl in doc.tables:
                        if tbl._tbl is prev:
                            celda = tbl.cell(0, 0)
                            # Limpiar contenido previo
                            celda.text = ""
                            p = celda.paragraphs[0]
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            run = p.add_run()
                            run.add_picture(str(ruta_imagen), width=Inches(5.8))
                            doc.save(str(RUTA_DOC))
                            print(f"✓ Imagen insertada en Figura {num_figura}: {ruta_imagen.name}")
                            return True
                    break
            print(f"✗ No se encontró tabla placeholder antes de 'Figura {num_figura}.'")
            return False
    print(f"✗ No se encontró el texto 'Figura {num_figura}.' en el documento.")
    return False


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Uso: python insertar_imagen.py <num_figura> <ruta_imagen>")
        sys.exit(1)
    num = int(sys.argv[1])
    ruta = Path(sys.argv[2])
    if not ruta.exists():
        print(f"✗ La imagen {ruta} no existe.")
        sys.exit(1)
    ok = insertar_en_figura(num, ruta)
    sys.exit(0 if ok else 1)
